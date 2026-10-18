"""
Generate Deliverable 05: UI/UX Wireframe & Design System Catalog (Baltazar & Capoy)
Outputs: DOCX & PDF in deliverables/05_Baltazar_Capoy_UI_UX_Wireframes_Catalog/
"""

import os
from docx.shared import Inches, Pt
from doc_builder import (
    create_base_docx, add_heading_1, add_heading_2, add_heading_3,
    add_paragraph, add_bullet, add_callout, add_table,
    build_html_wrapper, html_to_pdf
)

OUT_DIR = os.path.abspath("deliverables/05_Baltazar_Capoy_UI_UX_Wireframes_Catalog")
os.makedirs(OUT_DIR, exist_ok=True)

WIREFRAMES = [
    ("01_Sign_In.png", "01. Authentication & Sign-In Portal", "Initial entrypoint featuring dual-role credentials, institutional branding, quick demo login helper pills, and CSRF protection."),
    ("02_Student_Profile.png", "02. Student Profile & Credentials", "Comprehensive student biodata display, program/year level verification, and self-service password update form."),
    ("03_Student_Grades.png", "03. Academic Term Grades & GPA Engine", "Term-by-term enrolled subject grade breakdown with real-time automated semester GPA calculation."),
    ("04_Student_Subjects.png", "04. Enrolled Subjects & Course Schedules", "Complete class roster directory featuring course codes, descriptions, units, professors, and classroom assignments."),
    ("05_Student_Document_Requests.png", "05. Student Document Request History", "Self-service tracking portal showing document request histories, pricing, copy counts, and live fulfillment statuses."),
    ("06_Student_New_Request_Modal.png", "06. New Document Request Modal", "Interactive submission modal with dynamic price computation (₱50 - ₱150) based on document type and copy quantity."),
    ("07_Registrar_Enrollment.png", "07. Registrar Enrollment Management", "Administrative student-to-course registration directory with academic term and school year filters."),
    ("08_Registrar_Enroll_Modal.png", "08. Course Enrollment Modal", "Enrollment dialogue allowing registrars to register students into curriculum courses with instant validation."),
    ("09_Registrar_Grades_Encoding.png", "09. Registrar Grades Encoding Grid", "Live grade encoding spreadsheet grid with instantaneous database persistence across grading scales (1.00 - 5.00, INC)."),
    ("10_Cashier_Payments.png", "10. Cashier Payment Ledger", "Centralized counter payment ledger displaying receipt numbers, student identifiers, amounts paid, and payment statuses."),
    ("11_Cashier_Record_Payment_Modal.png", "11. Record Counter Payment Modal", "Modal for processing tuition and fee payments with automated Official Receipt (OR) sequential number generation."),
    ("12_Cashier_Official_Receipt.png", "12. University Official Receipt (Printable)", "Official institutional receipt layout formatted for standard 8.5x11 printing with cashier signature and watermark."),
    ("13_Department_Clearance.png", "13. Department Clearance Evaluation", "Multi-department clearance grid enabling academic chairs to toggle student clearances between Cleared, Hold, and Pending."),
    ("14_Admin_User_Management.png", "14. Administrative User Account Management", "SuperAdmin user directory with account creation, role assignment (Student, Registrar, Cashier, DeptStaff, Admin), and status controls."),
    ("15_Admin_Document_Requests.png", "15. Admin Document Fulfillment Queue", "Administrative queue enabling staff to advance document request lifecycles from Pending to Processing, Ready for Release, and Released.")
]

def generate_docx():
    doc = create_base_docx(
        title="UI/UX Design System Specification\n& Wireframe Catalog",
        subtitle="15 High-Fidelity Interface Wireframes, Design Tokens, and User Flow Walkthroughs",
        author="Baltazar & Capoy",
        role="Developers / UI-UX Designers",
        deliverable_code="TASK-3B (Baltazar & Capoy)"
    )

    add_heading_1(doc, "1. Design Philosophy & Institutional Branding")
    add_paragraph(doc,
        "The user experience of the CuyoTech Student Services Information System (SSIS) was designed around "
        "three primary design tenets: Institutional Authority, Cognitive Simplicity, and Zero-Friction Accessibility. "
        "Every view reflects CuyoTech University's official brand identity while prioritizing high clarity for "
        "dense academic tables and transaction modals."
    )

    headers_tokens = ["Token Category", "Design Token", "CSS Hex / Value", "Usage & Rationale"]
    rows_tokens = [
        ["Primary Brand", "--primary-navy", "#1e3a8a", "University authority; used in top navigation, active links, primary buttons, and table headers."],
        ["Secondary Accent", "--accent-gold", "#d97706 / #eab308", "Academic prestige; used in active tabs, badges, status highlights, and focus borders."],
        ["Neutral Surface", "--bg-canvas", "#f8fafc", "Light, clean background minimizing eye strain during extended administrative sessions."],
        ["Dark Typography", "--text-main", "#0f172a", "High-contrast slate black exceeding WCAG AAA contrast ratio standards for text readability."],
        ["Success State", "--badge-success", "#dcfce7 / #15803d", "Clear visual indicator for Cleared clearance, Paid fees, and Released documents."],
        ["Alert / Hold State", "--badge-danger", "#fee2e2 / #b91c1c", "High-visibility warning indicator for Disciplinary/Financial Holds and Deactivated accounts."]
    ]
    add_table(doc, headers_tokens, rows_tokens, [Inches(1.3), Inches(1.5), Inches(1.5), Inches(2.2)])

    add_heading_1(doc, "2. High-Fidelity Wireframes Catalog (15 System Views)")
    add_paragraph(doc, 
        "Below is the comprehensive catalog of all fifteen (15) functional user interface wireframes "
        "implemented within the SSIS application. Each screen includes design rationale, primary user actions, "
        "and embedded interface captures."
    )

    wf_base = os.path.abspath("assets/figma_wireframes")
    for img_name, screen_title, description in WIREFRAMES:
        add_heading_2(doc, screen_title)
        add_paragraph(doc, description)
        img_path = os.path.join(wf_base, img_name)
        if os.path.exists(img_path):
            try:
                p_img = doc.add_paragraph()
                p_img.add_run().add_picture(img_path, width=Inches(6.2))
            except Exception as e:
                add_paragraph(doc, f"[Image: {img_name}]")
        doc.add_paragraph()

    docx_path = os.path.join(OUT_DIR, "UI_UX_Wireframes_Catalog.docx")
    doc.save(docx_path)
    print(f"Generated DOCX: {docx_path}")
    return docx_path

def generate_pdf():
    wf_base = os.path.abspath("assets/figma_wireframes")
    screens_html = ""
    for img_name, screen_title, description in WIREFRAMES:
        img_path = os.path.join(wf_base, img_name)
        screens_html += f"""
<h2>{screen_title}</h2>
<p>{description}</p>
<div>
  <img src="file://{img_path}" class="wireframe-img" alt="{screen_title}">
</div>
<hr style="border:none;border-top:1px solid #e2e8f0;margin:24px 0;">
"""

    html_content = f"""
<h1>1. Design Philosophy & Institutional Branding</h1>
<p>
The user experience of the CuyoTech Student Services Information System (SSIS) was designed around
three primary design tenets: Institutional Authority, Cognitive Simplicity, and Zero-Friction Accessibility.
Every view reflects CuyoTech University's official brand identity while prioritizing high clarity for
dense academic tables and transaction modals.
</p>

<table class="data-table">
  <thead>
    <tr>
      <th>Token Category</th>
      <th>Design Token</th>
      <th>CSS Hex / Value</th>
      <th>Usage & Rationale</th>
    </tr>
  </thead>
  <tbody>
    <tr><td><strong>Primary Brand</strong></td><td>--primary-navy</td><td>#1e3a8a</td><td>University authority; used in top navigation, active links, primary buttons.</td></tr>
    <tr><td><strong>Secondary Accent</strong></td><td>--accent-gold</td><td>#d97706 / #eab308</td><td>Academic prestige; used in active tabs, badges, status highlights.</td></tr>
    <tr><td><strong>Neutral Surface</strong></td><td>--bg-canvas</td><td>#f8fafc</td><td>Light, clean background minimizing eye strain during extended administrative sessions.</td></tr>
    <tr><td><strong>Dark Typography</strong></td><td>--text-main</td><td>#0f172a</td><td>High-contrast slate black exceeding WCAG AAA contrast ratio standards.</td></tr>
    <tr><td><strong>Success State</strong></td><td>--badge-success</td><td>#dcfce7 / #15803d</td><td>Clear visual indicator for Cleared clearance, Paid fees, and Released documents.</td></tr>
    <tr><td><strong>Alert / Hold State</strong></td><td>--badge-danger</td><td>#fee2e2 / #b91c1c</td><td>High-visibility warning indicator for Disciplinary/Financial Holds.</td></tr>
  </tbody>
</table>

<h1>2. High-Fidelity Wireframes Catalog (15 System Views)</h1>
<p>
Below is the comprehensive catalog of all fifteen (15) functional user interface wireframes
implemented within the SSIS application:
</p>

{screens_html}
"""
    html_doc = build_html_wrapper(
        title="UI/UX Design System Specification & Wireframe Catalog",
        deliverable_code="TASK-3B (Baltazar & Capoy - Developers/Designers)",
        author_role="Baltazar & Capoy (Developers / UI-UX Designers)",
        content_html=html_content
    )
    html_path = os.path.join(OUT_DIR, "UI_UX_Wireframes_Catalog.html")
    pdf_path = os.path.join(OUT_DIR, "UI_UX_Wireframes_Catalog.pdf")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_doc)
    html_to_pdf(html_path, pdf_path)
    print(f"Generated PDF: {pdf_path}")
    return pdf_path

if __name__ == "__main__":
    generate_docx()
    generate_pdf()
