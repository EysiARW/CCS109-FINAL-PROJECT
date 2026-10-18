"""
Document Builder Utility for CuyoTech SSIS Case Study Deliverables
Generates professional DOCX and PDF files with unified branding.
"""

import os
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

CHROME_PATH = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

# Color Palette: CuyoTech Navy & Gold
NAVY_HEX = "#1e3a8a"
GOLD_HEX = "#d97706"
DARK_HEX = "#0f172a"
LIGHT_BG_HEX = "#f8fafc"
BORDER_HEX = "#cbd5e1"

COLOR_NAVY = RGBColor(30, 58, 138)
COLOR_GOLD = RGBColor(217, 119, 6)
COLOR_DARK = RGBColor(15, 23, 42)
COLOR_GRAY = RGBColor(100, 116, 139)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex.replace('#', ''))
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for name, w in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{name}')
        node.set(qn('w:w'), str(w))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_base_docx(title, subtitle, author, role, deliverable_code):
    doc = Document()
    
    # Page Margins: 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"CuyoTech SSIS | {deliverable_code}")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_GRAY
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("CCS109 Case Study - Student Services Information System | Confidential & Academic Use")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8)
        frun.font.color.rgb = COLOR_GRAY

    # Title Page Block
    logo_path = os.path.abspath("CuyoTech_University_Logo_sm.png")
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.add_run().add_picture(logo_path, width=Inches(1.8))
    
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("CUYOTECH UNIVERSITY\nCollege of Computing Studies\nDepartment of Information Technology")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(11)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_GRAY

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(18)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(title)
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = COLOR_GOLD

    doc.add_paragraph()

    # Meta Table Box
    tbl = doc.add_table(rows=5, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Course / Subject:", "CCS109 - Case Study / Software Engineering"),
        ("System:", "Student Services Information System (SSIS)"),
        ("Assigned Lead:", f"{author} ({role})"),
        ("Project Team:", "Bautista (Lead), Aquino (Scrum Master), Baltazar & Capoy (Devs), Balbero (QA)"),
        ("Submission Date:", "October 2026 | Academic Year 2026–2027")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = tbl.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, LIGHT_BG_HEX)
        set_cell_background(c1, "#ffffff")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_NAVY
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK

    doc.add_page_break()
    return doc

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = COLOR_NAVY
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = COLOR_GOLD
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = COLOR_NAVY
    return p

def add_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_DARK
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_NAVY
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK
    return p

def add_callout(doc, text, title="NOTE / CRITICAL REQUIREMENT"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.rows[0].cells[0]
    c.width = Inches(6.5)
    set_cell_background(c, LIGHT_BG_HEX)
    set_cell_margins(c, 120, 120, 150, 150)
    
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r_t = p.add_run(f"📌 {title}: ")
    r_t.font.name = "Calibri"
    r_t.font.bold = True
    r_t.font.size = Pt(9.5)
    r_t.font.color.rgb = COLOR_NAVY
    
    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = COLOR_DARK

def add_table(doc, headers, rows, col_widths=None):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header Row
    hdr_row = tbl.rows[0]
    for idx, h_text in enumerate(headers):
        c = hdr_row.cells[idx]
        if col_widths and idx < len(col_widths):
            c.width = col_widths[idx]
        set_cell_background(c, NAVY_HEX)
        set_cell_margins(c, 80, 80, 100, 100)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h_text)
        r.font.name = "Calibri"
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    # Data Rows
    for r_idx, r_data in enumerate(rows):
        row = tbl.rows[r_idx + 1]
        bg_col = LIGHT_BG_HEX if r_idx % 2 == 1 else "#ffffff"
        for c_idx, val in enumerate(r_data):
            c = row.cells[c_idx]
            if col_widths and c_idx < len(col_widths):
                c.width = col_widths[c_idx]
            set_cell_background(c, bg_col)
            set_cell_margins(c, 60, 60, 100, 100)
            p = c.paragraphs[0]
            r = p.add_run(str(val))
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_DARK
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return tbl

def html_to_pdf(html_path, pdf_path):
    cmd = [
        CHROME_PATH,
        "--headless",
        "--disable-gpu",
        "--no-pdf-header-footer",
        f"--print-to-pdf={pdf_path}",
        f"file://{os.path.abspath(html_path)}"
    ]
    subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    return os.path.exists(pdf_path)

def build_html_wrapper(title, deliverable_code, author_role, content_html):
    logo_abs = os.path.abspath("CuyoTech_University_Logo_sm.png")
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
@page {{
  size: A4 portrait;
  margin: 18mm 18mm 20mm 18mm;
}}
body {{
  font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
  color: #0f172a;
  background: #ffffff;
  line-height: 1.55;
  font-size: 10pt;
  margin: 0;
}}
.cover-page {{
  text-align: center;
  padding-top: 40px;
  page-break-after: always;
}}
.cover-logo {{
  width: 140px;
  margin-bottom: 20px;
}}
.inst-title {{
  font-size: 13pt;
  font-weight: 700;
  color: #475569;
  letter-spacing: 0.5px;
  margin: 0 0 40px 0;
  text-transform: uppercase;
  line-height: 1.4;
}}
.doc-title {{
  font-size: 24pt;
  font-weight: 800;
  color: #1e3a8a;
  margin: 0 0 10px 0;
  line-height: 1.25;
}}
.doc-subtitle {{
  font-size: 13pt;
  font-weight: 600;
  color: #d97706;
  margin: 0 0 40px 0;
}}
.meta-box {{
  width: 85%;
  margin: 0 auto;
  border: 1px solid #cbd5e1;
  border-radius: 6px;
  overflow: hidden;
  text-align: left;
  border-collapse: collapse;
}}
.meta-box td {{
  padding: 8px 14px;
  font-size: 9.5pt;
  border-bottom: 1px solid #e2e8f0;
}}
.meta-label {{
  background: #f8fafc;
  font-weight: 700;
  color: #1e3a8a;
  width: 32%;
}}
.meta-val {{
  color: #1e293b;
}}

/* Typography */
h1 {{
  color: #1e3a8a;
  font-size: 15pt;
  border-bottom: 2px solid #e2e8f0;
  padding-bottom: 5px;
  margin-top: 26px;
  margin-bottom: 10px;
  page-break-after: avoid;
}}
h2 {{
  color: #b45309;
  font-size: 12pt;
  margin-top: 18px;
  margin-bottom: 6px;
  page-break-after: avoid;
}}
h3 {{
  color: #1e3a8a;
  font-size: 10.5pt;
  margin-top: 14px;
  margin-bottom: 4px;
  page-break-after: avoid;
}}
p {{
  margin: 0 0 8px 0;
}}
ul, ol {{
  margin: 0 0 10px 0;
  padding-left: 20px;
}}
li {{
  margin-bottom: 4px;
}}
strong {{
  color: #0f172a;
}}
table.data-table {{
  width: 100%;
  border-collapse: collapse;
  margin: 12px 0 16px 0;
  font-size: 8.5pt;
  page-break-inside: auto;
}}
table.data-table tr {{
  page-break-inside: avoid;
  page-break-after: auto;
}}
table.data-table th {{
  background: #1e3a8a;
  color: #ffffff;
  text-align: left;
  padding: 7px 10px;
  font-weight: 700;
  border: 1px solid #1e3a8a;
}}
table.data-table td {{
  padding: 6px 10px;
  border: 1px solid #cbd5e1;
  color: #1e293b;
}}
table.data-table tr:nth-child(even) td {{
  background: #f8fafc;
}}
.callout {{
  background: #eff6ff;
  border-left: 4px solid #1e3a8a;
  padding: 10px 14px;
  margin: 12px 0;
  border-radius: 0 6px 6px 0;
  font-size: 9pt;
}}
.callout-title {{
  font-weight: 700;
  color: #1e3a8a;
  margin-bottom: 4px;
}}
.badge {{
  display: inline-block;
  padding: 2px 7px;
  border-radius: 4px;
  font-size: 8pt;
  font-weight: 700;
}}
.badge-pass {{ background: #dcfce7; color: #15803d; }}
.badge-must {{ background: #fee2e2; color: #b91c1c; }}
.badge-should {{ background: #fef3c7; color: #b45309; }}
.badge-could {{ background: #e0e7ff; color: #4338ca; }}
.badge-done {{ background: #dcfce7; color: #166534; }}
.code-box {{
  background: #0f172a;
  color: #f8fafc;
  padding: 10px 14px;
  border-radius: 6px;
  font-family: Consolas, "Liberation Mono", Menlo, monospace;
  font-size: 8pt;
  line-height: 1.4;
  white-space: pre-wrap;
  margin: 10px 0;
}}
.gantt-bar {{
  background: #1e3a8a;
  height: 12px;
  border-radius: 3px;
}}
.wireframe-img {{
  max-width: 95%;
  height: auto;
  border: 1px solid #cbd5e1;
  border-radius: 6px;
  box-shadow: 0 2px 6px rgba(0,0,0,0.08);
  margin: 8px 0;
}}
.footer-text {{
  font-size: 8pt;
  color: #64748b;
  text-align: center;
  border-top: 1px solid #e2e8f0;
  padding-top: 6px;
  margin-top: 30px;
}}
</style>
</head>
<body>

<div class="cover-page">
  <img src="file://{logo_abs}" class="cover-logo" alt="Logo">
  <div class="inst-title">CuyoTech University<br>College of Computing Studies<br>Department of Information Technology</div>
  <div class="doc-title">{title}</div>
  <div class="doc-subtitle">Case Study Deliverable: {deliverable_code}</div>
  
  <table class="meta-box">
    <tr><td class="meta-label">Course / Subject</td><td class="meta-val">CCS109 - Case Study / Software Engineering</td></tr>
    <tr><td class="meta-label">System Focus</td><td class="meta-val">Student Services Information System (SSIS)</td></tr>
    <tr><td class="meta-label">Assigned Lead</td><td class="meta-val">{author_role}</td></tr>
    <tr><td class="meta-label">Project Team</td><td class="meta-val">Bautista (Lead), Aquino (Scrum Master), Baltazar & Capoy (Devs), Balbero (QA)</td></tr>
    <tr><td class="meta-label">Academic Year / Term</td><td class="meta-val">Academic Year 2026–2027 | 1st Semester</td></tr>
  </table>
</div>

<div class="main-content">
{content_html}
</div>

<div class="footer-text">
  CuyoTech SSIS Case Study Deliverable | College of Computing Studies | Confidential & Academic Use
</div>

</body>
</html>"""
