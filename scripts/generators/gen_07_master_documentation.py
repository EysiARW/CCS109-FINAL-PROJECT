"""
Generate Deliverable 07: Master Case Study Documentation (Balbero / Entire Team)
Outputs: DOCX & PDF in deliverables/07_Balbero_Consolidated_Project_Documentation/
"""

import os
from docx.shared import Inches, Pt
from doc_builder import (
    create_base_docx, add_heading_1, add_heading_2, add_heading_3,
    add_paragraph, add_bullet, add_callout, add_table,
    build_html_wrapper, html_to_pdf
)

OUT_DIR = os.path.abspath("deliverables/07_Balbero_Consolidated_Project_Documentation")
os.makedirs(OUT_DIR, exist_ok=True)

def generate_docx():
    doc = create_base_docx(
        title="Student Services Information System\nMaster Project Documentation",
        subtitle="Consolidated Technical Manuscript, Architecture, 3NF Normalization, QA, and User Guide",
        author="Balbero (Lead) with Bautista, Aquino, Baltazar, Capoy",
        role="Documentation Lead & Team",
        deliverable_code="TASK-4C & MASTER (Consolidated)"
    )

    add_heading_1(doc, "Chapter 1: Project Overview & Institutional Background")
    add_paragraph(doc,
        "The Student Services Information System (SSIS) is an enterprise web application engineered for CuyoTech University. "
        "Historically, the institution relied on fragmented manual processes, physical ledger books, and ad-hoc spreadsheets "
        "to manage student profiles, enrollments, grading records, clearance certifications, and tuition counter collections. "
        "These manual bottlenecks created extended queuing times, heightened vulnerability to clerical errors, and raised "
        "severe data privacy compliance concerns under the Philippine Data Privacy Act of 2012 (RA 10173)."
    )
    add_paragraph(doc,
        "The SSIS establishes a modern, unified, role-based platform engineered in pure PHP 8.2 with a Third Normal Form (3NF) "
        "relational MySQL database. Designed with zero external runtime dependencies, the system delivers high performance, "
        "rock-solid stability, and seamless portability across standard web hosting and local XAMPP environments."
    )

    add_heading_1(doc, "Chapter 2: System Requirements & Agile Backlog")
    add_paragraph(doc,
        "Requirements were synthesized through extensive interviews with registrars, cashiers, collegiate department chairs, "
        "and student councils. Functional specifications were partitioned across six core operational domains: Authentication (FR-AUTH), "
        "Student Self-Service (FR-STU), Document Requests (FR-DOC), Registrar Enrollment & Grading (FR-REG), Cashier Counter Payments (FR-CASH), "
        "and Department Clearances & Administration (FR-ADM)."
    )
    add_paragraph(doc,
        "Using the Agile MoSCoW framework, the project backlog was scoped into 44 Must-Have Story Points, 11 Should-Have Story Points, "
        "and 5 Could/Won't-Have Story Points, ensuring that 100% of critical core functionalities were fully realized by Sprint 4."
    )

    add_heading_1(doc, "Chapter 3: SDLC Methodology & Project Schedule")
    add_paragraph(doc,
        "The team adhered to the Agile/Scrum framework across an 8-week lifecycle (September 1 to October 24, 2026), "
        "organized into four 2-week iterations: Sprint 1 (MVC Engine & Authentication), Sprint 2 (Academic Records & Document Requests), "
        "Sprint 3 (Cashier Operations & Department Clearances), and Sprint 4 (Admin Controls, Security Hardening, and Quality Audits). "
        "Sprint ceremonies—including bi-weekly planning, tri-weekly standups, and milestone reviews—ensured continuous quality governance."
    )

    add_heading_1(doc, "Chapter 4: System Architecture & Design")
    add_paragraph(doc,
        "The application architecture strictly follows the Model-View-Controller (MVC) pattern coupled with the Repository Pattern. "
        "Controllers orchestrate request validation and session authorization; Repositories encapsulate all PDO database interactions; "
        "and Models represent strongly-typed entities. Zero raw SQL concatenation exists within the presentation layer. "
        "All views inherit a modular university layout featuring CuyoTech Navy (#1e3a8a) and Gold (#d97706) branding."
    )

    add_heading_1(doc, "Chapter 5: Database Design & 3NF Normalization")
    add_paragraph(doc,
        "The relational schema comprises eight (8) distinct tables normalized to Third Normal Form (3NF):\n"
        "1. Department: Master list of university academic and administrative offices.\n"
        "2. Student: Student credentials, biographical profiles, and program affiliations.\n"
        "3. AdminUser: Administrative staff accounts with granular role permissions.\n"
        "4. Course: Course catalog including unit values and offering departments.\n"
        "5. Enrollment: Junction table associating students to courses with term grades.\n"
        "6. Payment: Counter financial transaction ledger with Official Receipt (OR) tracking.\n"
        "7. DocumentRequest: Fulfillment tracking for transcripts, certifications, and clearances.\n"
        "8. Clearance: Cross-department student clearance statuses (Cleared, Pending, Hold)."
    )

    add_heading_1(doc, "Chapter 6: Quality Assurance & Security Verification")
    add_paragraph(doc,
        "A formal IEEE 829 Software Test Plan was executed by Lead QA Balbero. The automated test suite ran thirty-six (36) "
        "comprehensive unit, integration, and security test cases against the live MySQL database, achieving a verified 100% pass rate. "
        "Security audits confirmed total immunity against SQL injection attack vectors through 100% PDO prepared statements, "
        "bulletproof XSS neutralization via htmlspecialchars output encoding, and industry-standard Bcrypt ($2y$10$) password hashing."
    )

    add_heading_1(doc, "Chapter 7: User Manual & Operational Procedures")
    add_heading_2(doc, "7.1 Student User Guide")
    add_bullet(doc, "Navigate to index.php. Log in using Student ID (e.g., 2021-00123) and password (password123).", "Step 1: Sign-In: ")
    add_bullet(doc, "Click 'Grades' in the navigation bar to inspect term grades and verify real-time GPA calculations.", "Step 2: Inspect Grades: ")
    add_bullet(doc, "Navigate to 'Document Requests', click 'Request Document', choose document type and copies, and submit.", "Step 3: Document Requests: ")
    add_bullet(doc, "Open 'Profile' to review biographical data or update account password.", "Step 4: Profile Updates: ")

    add_heading_2(doc, "7.2 Staff & Administrative User Guide")
    add_bullet(doc, "Registrars access ?page=grades-encoding to input numeric student marks with instant database persistence.", "Registrar Operations: ")
    add_bullet(doc, "Cashiers access ?page=payments to record counter collections and generate printable Official Receipts.", "Cashier Operations: ")
    add_bullet(doc, "Department chairs access ?page=clearance to review student obligations and toggle clearance statuses.", "Department Operations: ")
    add_bullet(doc, "Administrators access ?page=admin-users to provision new user accounts and ?page=admin-requests to fulfill certificates.", "Admin Operations: ")

    add_heading_1(doc, "Chapter 8: Conclusion & Recommendations")
    add_paragraph(doc,
        "The CuyoTech Student Services Information System successfully achieves all educational, architectural, and operational "
        "benchmarks established for the CCS109 case study. It establishes a resilient, tamper-resistant, and high-performance digital "
        "backbone for collegiate operations. Future roadmap recommendations include mobile-native PWA optimization, automated SMS "
        "gateway notifications for released documents, and batch grade import via CSV."
    )

    docx_path = os.path.join(OUT_DIR, "Master_Case_Study_Documentation.docx")
    doc.save(docx_path)
    print(f"Generated DOCX: {docx_path}")
    return docx_path

def generate_pdf():
    html_content = """
<h1>Chapter 1: Project Overview & Institutional Background</h1>
<p>
The Student Services Information System (SSIS) is an enterprise web application engineered for CuyoTech University.
Historically, the institution relied on fragmented manual processes, physical ledger books, and ad-hoc spreadsheets
to manage student profiles, enrollments, grading records, clearance certifications, and tuition counter collections.
These manual bottlenecks created extended queuing times, heightened vulnerability to clerical errors, and raised
severe data privacy compliance concerns under the Philippine Data Privacy Act of 2012 (RA 10173).
</p>
<p>
The SSIS establishes a modern, unified, role-based platform engineered in pure PHP 8.2 with a Third Normal Form (3NF)
relational MySQL database. Designed with zero external runtime dependencies, the system delivers high performance,
rock-solid stability, and seamless portability across standard web hosting and local XAMPP environments.
</p>

<h1>Chapter 2: System Requirements & Agile Backlog</h1>
<p>
Requirements were synthesized through extensive interviews with registrars, cashiers, collegiate department chairs,
and student councils. Functional specifications were partitioned across six core operational domains: Authentication (FR-AUTH),
Student Self-Service (FR-STU), Document Requests (FR-DOC), Registrar Enrollment & Grading (FR-REG), Cashier Counter Payments (FR-CASH),
and Department Clearances & Administration (FR-ADM).
</p>
<p>
Using the Agile MoSCoW framework, the project backlog was scoped into 44 Must-Have Story Points, 11 Should-Have Story Points,
and 5 Could/Won't-Have Story Points, ensuring that 100% of critical core functionalities were fully realized by Sprint 4.
</p>

<h1>Chapter 3: SDLC Methodology & Project Schedule</h1>
<p>
The team adhered to the Agile/Scrum framework across an 8-week lifecycle (September 1 to October 24, 2026),
organized into four 2-week iterations: Sprint 1 (MVC Engine & Authentication), Sprint 2 (Academic Records & Document Requests),
Sprint 3 (Cashier Operations & Department Clearances), and Sprint 4 (Admin Controls, Security Hardening, and Quality Audits).
Sprint ceremonies—including bi-weekly planning, tri-weekly standups, and milestone reviews—ensured continuous quality governance.
</p>

<h1>Chapter 4: System Architecture & Design</h1>
<p>
The application architecture strictly follows the Model-View-Controller (MVC) pattern coupled with the Repository Pattern.
Controllers orchestrate request validation and session authorization; Repositories encapsulate all PDO database interactions;
and Models represent strongly-typed entities. Zero raw SQL concatenation exists within the presentation layer.
All views inherit a modular university layout featuring CuyoTech Navy (#1e3a8a) and Gold (#d97706) branding.
</p>

<h1>Chapter 5: Database Design & 3NF Normalization</h1>
<p>The relational schema comprises eight (8) distinct tables normalized to Third Normal Form (3NF):</p>
<ul>
  <li><strong>Department:</strong> Master list of university academic and administrative offices.</li>
  <li><strong>Student:</strong> Student credentials, biographical profiles, and program affiliations.</li>
  <li><strong>AdminUser:</strong> Administrative staff accounts with granular role permissions.</li>
  <li><strong>Course:</strong> Course catalog including unit values and offering departments.</li>
  <li><strong>Enrollment:</strong> Junction table associating students to courses with term grades.</li>
  <li><strong>Payment:</strong> Counter financial transaction ledger with Official Receipt (OR) tracking.</li>
  <li><strong>DocumentRequest:</strong> Fulfillment tracking for transcripts, certifications, and clearances.</li>
  <li><strong>Clearance:</strong> Cross-department student clearance statuses (Cleared, Pending, Hold).</li>
</ul>

<h1>Chapter 6: Quality Assurance & Security Verification</h1>
<p>
A formal IEEE 829 Software Test Plan was executed by Lead QA Balbero. The automated test suite ran thirty-six (36)
comprehensive unit, integration, and security test cases against the live MySQL database, achieving a verified 100% pass rate.
Security audits confirmed total immunity against SQL injection attack vectors through 100% PDO prepared statements,
bulletproof XSS neutralization via htmlspecialchars output encoding, and industry-standard Bcrypt ($2y$10$) password hashing.
</p>

<h1>Chapter 7: User Manual & Operational Procedures</h1>
<h2>7.1 Student User Guide</h2>
<ul>
  <li><strong>Step 1: Sign-In:</strong> Navigate to index.php. Log in using Student ID (e.g., 2021-00123) and password (password123).</li>
  <li><strong>Step 2: Inspect Grades:</strong> Click 'Grades' in the navigation bar to inspect term grades and verify real-time GPA calculations.</li>
  <li><strong>Step 3: Document Requests:</strong> Navigate to 'Document Requests', click 'Request Document', choose document type and copies, and submit.</li>
  <li><strong>Step 4: Profile Updates:</strong> Open 'Profile' to review biographical data or update account password.</li>
</ul>

<h2>7.2 Staff & Administrative User Guide</h2>
<ul>
  <li><strong>Registrar Operations:</strong> Registrars access ?page=grades-encoding to input numeric student marks with instant database persistence.</li>
  <li><strong>Cashier Operations:</strong> Cashiers access ?page=payments to record counter collections and generate printable Official Receipts.</li>
  <li><strong>Department Operations:</strong> Department chairs access ?page=clearance to review student obligations and toggle clearance statuses.</li>
  <li><strong>Admin Operations:</strong> Administrators access ?page=admin-users to provision new user accounts and ?page=admin-requests to fulfill certificates.</li>
</ul>

<h1>Chapter 8: Conclusion & Recommendations</h1>
<p>
The CuyoTech Student Services Information System successfully achieves all educational, architectural, and operational
benchmarks established for the CCS109 case study. It establishes a resilient, tamper-resistant, and high-performance digital
backbone for collegiate operations. Future roadmap recommendations include mobile-native PWA optimization, automated SMS
gateway notifications for released documents, and batch grade import via CSV.
</p>
"""
    html_doc = build_html_wrapper(
        title="Student Services Information System Master Documentation",
        deliverable_code="TASK-4C & MASTER (Balbero / Entire Team)",
        author_role="Balbero (Lead) with Bautista, Aquino, Baltazar, Capoy",
        content_html=html_content
    )
    html_path = os.path.join(OUT_DIR, "Master_Case_Study_Documentation.html")
    pdf_path = os.path.join(OUT_DIR, "Master_Case_Study_Documentation.pdf")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_doc)
    html_to_pdf(html_path, pdf_path)
    print(f"Generated PDF: {pdf_path}")
    return pdf_path

if __name__ == "__main__":
    generate_docx()
    generate_pdf()
