"""
Generate Deliverable 04: Ethical Analysis & Final Review (Bautista & Aquino)
Outputs: DOCX & PDF in deliverables/04_Bautista_Aquino_Ethical_Analysis_and_Final_Review/
"""

import os
from docx.shared import Inches, Pt
from doc_builder import (
    create_base_docx, add_heading_1, add_heading_2, add_heading_3,
    add_paragraph, add_bullet, add_callout, add_table,
    build_html_wrapper, html_to_pdf
)

OUT_DIR = os.path.abspath("deliverables/04_Bautista_Aquino_Ethical_Analysis_and_Final_Review")
os.makedirs(OUT_DIR, exist_ok=True)

def generate_docx():
    doc = create_base_docx(
        title="Ethical Analysis, Legal Compliance\n& Pre-Submission Quality Review",
        subtitle="Philippine Data Privacy Act (RA 10173), Segregation of Duties, and Quality Gate Sign-Off",
        author="Bautista & Aquino",
        role="Team Leader & SCRUM Master",
        deliverable_code="TASK-1C, 1D & 2D (Bautista & Aquino)"
    )

    add_heading_1(doc, "1. Ethical Principles in Educational Computing")
    add_paragraph(doc,
        "Student Information Systems manage highly sensitive academic and personally identifiable records. "
        "A breach of integrity or unauthorized exposure of academic grades, financial standings, or disciplinary holds "
        "can irreversibly damage a student's career prospects, privacy, and psychological well-being. "
        "Consequently, the architectural design of CuyoTech SSIS was governed by strict ethical imperatives: "
        "fiduciary duty of care, confidentiality of academic marks, transparency in fee calculation, and algorithmic accountability."
    )

    add_heading_1(doc, "2. Legal Compliance: Philippine Data Privacy Act of 2012 (RA 10173)")
    add_paragraph(doc,
        "As an academic institutional system developed in the Philippines, the SSIS was rigorously audited against the "
        "statutory mandates of Republic Act No. 10173 (Data Privacy Act of 2012) and its Implementing Rules and Regulations (IRR):"
    )

    headers_dpa = ["DPA Statutory Pillar", "Legal Mandate Description", "SSIS Technical Enforcement"]
    rows_dpa = [
        ["Principle of Transparency", "Data subjects must know what data is collected, why, and how it is processed.", "Clear student profile views explicitly displaying collected fields: StudentID, full name, email, program, address."],
        ["Legitimate Purpose", "Processing must be strictly bounded by educational service delivery.", "Data is collected strictly to fulfill enrollment, grade transcription, graduation clearance, and official receipt issuance."],
        ["Proportionality", "Only the minimum necessary data points shall be acquired and stored.", "No superfluous biographical or intrusive metadata is acquired. Sensitive passwords stored exclusively as one-way Bcrypt hashes."],
        ["Right to Access", "Students have the legal right to inspect their complete educational records on demand.", "Dedicated Student Portal (?page=grades and ?page=subjects) available 24/7 with zero human gatekeeping."],
        ["Right to Rectification", "Inaccurate personal data must be promptly correctable by the data subject.", "Self-service password update (?page=profile) and administrative student profile editing workflows."],
        ["Security of Personal Data", "Obligation to implement technical safeguards against leaks and unauthorized access.", "PDO prepared statements (100% coverage), Bcrypt password hashing, session regeneration, and XSS output escaping."]
    ]
    add_table(doc, headers_dpa, rows_dpa, [Inches(1.5), Inches(2.2), Inches(2.8)])

    add_heading_1(doc, "3. Segregation of Duties (SoD) & Anti-Fraud Governance")
    add_paragraph(doc,
        "A critical vulnerability in manual campus systems is the consolidation of conflicting administrative powers. "
        "To prevent institutional fraud, bribery, and grade tampering, SSIS implements strict Segregation of Duties (SoD) "
        "enforced through Role-Based Access Control (RBAC):"
    )
    add_bullet(doc, "Registrars encode grades and process academic credentials but are completely blocked from collecting cash payments or altering financial records.", "Registrar vs. Cashier Demarcation: ")
    add_bullet(doc, "Cashiers record counter payments and issue official receipts (OR) but possess zero access to student course enrollments or grade rosters.", "Cashier Independence: ")
    add_bullet(doc, "Department Chairs and evaluators can update departmental clearance statuses (Cleared, Hold) but cannot issue academic transcripts.", "Clearance Segregation: ")
    add_bullet(doc, "All actions are attributed to authenticated accounts, establishing full non-repudiation across student records.", "Auditability & Traceability: ")

    add_heading_1(doc, "4. Requirements Traceability Matrix (RTM)")
    add_paragraph(doc, "The matrix below maps each requirements specification to its corresponding codebase implementation and test verification:")

    headers_rtm = ["Req ID", "Requirement Description", "Source Controller / View", "Database Table", "Automated QA Verification"]
    rows_rtm = [
        ["FR-AUTH-01", "Student Multi-Identifier Login", "AuthController.php / views/auth/login.php", "Student", "TC-AUTH-01 & TC-AUTH-02 (PASS)"],
        ["FR-AUTH-02", "Admin & Staff Role Authentication", "AuthController.php / views/auth/login.php", "AdminUser", "TC-AUTH-03 to TC-AUTH-06 (PASS)"],
        ["FR-STU-01", "Student Profile & Password Update", "ProfileController.php / views/student/profile.php", "Student", "TC-REPO-01 & TC-CRYPTO-01 (PASS)"],
        ["FR-STU-03", "Term Grades & Real-Time GPA Math", "AcademicController.php / views/student/grades.php", "Enrollment, Course", "TC-MATH-01 & TC-REPO-03 (PASS)"],
        ["FR-DOC-01", "Document Request & Fee Calculation", "RequestController.php / views/student/requests.php", "DocumentRequest", "TC-REPO-06 & TC-REPO-07 (PASS)"],
        ["FR-CASH-01", "Cashier Payment & Official Receipt", "CashierController.php / views/cashier/payments.php", "Payment", "TC-REPO-05 (PASS)"],
        ["FR-ADM-01", "Cross-Department Clearance Review", "DepartmentController.php / views/department/", "Clearance, Department", "TC-REPO-08 (PASS)"],
        ["NFR-SEC-01", "Bcrypt Password Cryptography", "AuthService.php / config/Database.php", "Student, AdminUser", "TC-CRYPTO-01 & TC-CRYPTO-02 (PASS)"],
        ["NFR-SEC-02", "SQL Injection Parameterization", "App\\Repositories\\* (All 7 Repositories)", "All 8 Tables", "TC-SQLI-01 to TC-SQLI-05 (PASS)"],
        ["NFR-SEC-03", "XSS Output Sanitization", "App\\Utils\\View.php & views/*", "N/A (Presentation)", "TC-XSS-01 (PASS)"]
    ]
    add_table(doc, headers_rtm, rows_rtm, [Inches(0.9), Inches(1.8), Inches(1.8), Inches(1.0), Inches(1.0)])

    add_heading_1(doc, "5. Pre-Submission Final Review Checklist")
    add_callout(doc,
        "This pre-submission quality gate was jointly audited by Team Leader Bautista and SCRUM Master Aquino. "
        "Every single technical, architectural, testing, and documentation requirement for CCS109 has been verified.",
        title="FINAL AUDIT QUALITY CERTIFICATION"
    )

    headers_check = ["Evaluation Criteria", "Required Standard", "Actual Implementation Status", "Verified By"]
    rows_check = [
        ["3NF Database Normalization", "Minimum 6 tables, no partial or transitive dependencies", "8 Relational Tables (1NF -> 2NF -> 3NF certified in DATABASE_DESIGN.md)", "Bautista (Lead)"],
        ["Zero-Dependency Architecture", "Pure PHP + MySQL, no Node/npm or external build steps", "Pure PHP 8.2 MVC runtime with automated offline demo fallback", "Aquino (Scrum Master)"],
        ["Code Syntax & Linting", "Zero fatal errors or unhandled exceptions across all scripts", "All 44 PHP files passed strict php -l syntax linting", "Balbero (QA)"],
        ["Software Test Plan Execution", "Formal test cases with verifiable execution evidence", "IEEE 829 test suite: 36 of 36 automated tests passed (100% success)", "Balbero (QA)"],
        ["Security Vulnerability Audit", "Zero high-risk SQLi or XSS vulnerabilities", "100% PDO prepared statements, Bcrypt password hashing verified", "Balbero (QA)"],
        ["UI/UX Wireframes Deliverable", "Visual wireframe catalog of all system views", "15 Wireframes captured & cataloged in figma_wireframes/", "Capoy & Baltazar"],
        ["System Presentation Deck", "Widescreen PowerPoint slides for case study defense", "18-slide professional PPTX presentation with institutional branding", "Aquino (Scrum Master)"]
    ]
    add_table(doc, headers_check, rows_check, [Inches(1.5), Inches(1.8), Inches(2.2), Inches(1.0)])

    docx_path = os.path.join(OUT_DIR, "Ethical_Analysis_and_Final_Review.docx")
    doc.save(docx_path)
    print(f"Generated DOCX: {docx_path}")
    return docx_path

def generate_pdf():
    html_content = """
<h1>1. Ethical Principles in Educational Computing</h1>
<p>
Student Information Systems manage highly sensitive academic and personally identifiable records.
A breach of integrity or unauthorized exposure of academic grades, financial standings, or disciplinary holds
can irreversibly damage a student's career prospects, privacy, and psychological well-being.
Consequently, the architectural design of CuyoTech SSIS was governed by strict ethical imperatives:
fiduciary duty of care, confidentiality of academic marks, transparency in fee calculation, and algorithmic accountability.
</p>

<h1>2. Legal Compliance: Philippine Data Privacy Act of 2012 (RA 10173)</h1>
<p>
As an academic institutional system developed in the Philippines, the SSIS was rigorously audited against the
statutory mandates of Republic Act No. 10173 (Data Privacy Act of 2012) and its Implementing Rules and Regulations (IRR):
</p>

<table class="data-table">
  <thead>
    <tr>
      <th>DPA Statutory Pillar</th>
      <th>Legal Mandate Description</th>
      <th>SSIS Technical Enforcement</th>
    </tr>
  </thead>
  <tbody>
    <tr>
      <td><strong>Principle of Transparency</strong></td>
      <td>Data subjects must know what data is collected, why, and how it is processed.</td>
      <td>Clear student profile views explicitly displaying collected fields: StudentID, full name, email, program, address.</td>
    </tr>
    <tr>
      <td><strong>Legitimate Purpose</strong></td>
      <td>Processing must be strictly bounded by educational service delivery.</td>
      <td>Data is collected strictly to fulfill enrollment, grade transcription, graduation clearance, and official receipt issuance.</td>
    </tr>
    <tr>
      <td><strong>Proportionality</strong></td>
      <td>Only the minimum necessary data points shall be acquired and stored.</td>
      <td>No superfluous biographical or intrusive metadata is acquired. Sensitive passwords stored exclusively as one-way Bcrypt hashes.</td>
    </tr>
    <tr>
      <td><strong>Right to Access</strong></td>
      <td>Students have the legal right to inspect their complete educational records on demand.</td>
      <td>Dedicated Student Portal (?page=grades and ?page=subjects) available 24/7 with zero human gatekeeping.</td>
    </tr>
    <tr>
      <td><strong>Right to Rectification</strong></td>
      <td>Inaccurate personal data must be promptly correctable by the data subject.</td>
      <td>Self-service password update (?page=profile) and administrative student profile editing workflows.</td>
    </tr>
    <tr>
      <td><strong>Security of Personal Data</strong></td>
      <td>Obligation to implement technical safeguards against leaks and unauthorized access.</td>
      <td>PDO prepared statements (100% coverage), Bcrypt password hashing, session regeneration, and XSS output escaping.</td>
    </tr>
  </tbody>
</table>

<h1>3. Segregation of Duties (SoD) & Anti-Fraud Governance</h1>
<p>
A critical vulnerability in manual campus systems is the consolidation of conflicting administrative powers.
To prevent institutional fraud, bribery, and grade tampering, SSIS implements strict Segregation of Duties (SoD)
enforced through Role-Based Access Control (RBAC):
</p>
<ul>
  <li><strong>Registrar vs. Cashier Demarcation:</strong> Registrars encode grades and process academic credentials but are completely blocked from collecting cash payments or altering financial records.</li>
  <li><strong>Cashier Independence:</strong> Cashiers record counter payments and issue official receipts (OR) but possess zero access to student course enrollments or grade rosters.</li>
  <li><strong>Clearance Segregation:</strong> Department Chairs and evaluators can update departmental clearance statuses (Cleared, Hold) but cannot issue academic transcripts.</li>
  <li><strong>Auditability & Traceability:</strong> All actions are attributed to authenticated accounts, establishing full non-repudiation across student records.</li>
</ul>

<h1>4. Requirements Traceability Matrix (RTM)</h1>
<table class="data-table">
  <thead>
    <tr>
      <th>Req ID</th>
      <th>Requirement Description</th>
      <th>Source Controller / View</th>
      <th>Database Table</th>
      <th>Automated QA Verification</th>
    </tr>
  </thead>
  <tbody>
    <tr><td>FR-AUTH-01</td><td>Student Multi-Identifier Login</td><td>AuthController.php / login.php</td><td>Student</td><td><span class="badge badge-pass">TC-AUTH-01 & 02 (PASS)</span></td></tr>
    <tr><td>FR-AUTH-02</td><td>Staff Role Authentication</td><td>AuthController.php / login.php</td><td>AdminUser</td><td><span class="badge badge-pass">TC-AUTH-03 to 06 (PASS)</span></td></tr>
    <tr><td>FR-STU-01</td><td>Student Profile & Password Update</td><td>ProfileController.php / profile.php</td><td>Student</td><td><span class="badge badge-pass">TC-REPO-01 (PASS)</span></td></tr>
    <tr><td>FR-STU-03</td><td>Term Grades & Real-Time GPA Math</td><td>AcademicController.php / grades.php</td><td>Enrollment, Course</td><td><span class="badge badge-pass">TC-MATH-01 (PASS)</span></td></tr>
    <tr><td>FR-DOC-01</td><td>Document Request & Fee Calculation</td><td>RequestController.php / requests.php</td><td>DocumentRequest</td><td><span class="badge badge-pass">TC-REPO-06 & 07 (PASS)</span></td></tr>
    <tr><td>FR-CASH-01</td><td>Cashier Payment & Official Receipt</td><td>CashierController.php / payments.php</td><td>Payment</td><td><span class="badge badge-pass">TC-REPO-05 (PASS)</span></td></tr>
    <tr><td>FR-ADM-01</td><td>Cross-Department Clearance Review</td><td>DepartmentController.php / clearance.php</td><td>Clearance, Department</td><td><span class="badge badge-pass">TC-REPO-08 (PASS)</span></td></tr>
    <tr><td>NFR-SEC-01</td><td>Bcrypt Password Cryptography</td><td>AuthService.php / Database.php</td><td>Student, AdminUser</td><td><span class="badge badge-pass">TC-CRYPTO-01 (PASS)</span></td></tr>
    <tr><td>NFR-SEC-02</td><td>SQL Injection Parameterization</td><td>App\\Repositories\\* (All 7 Repos)</td><td>All 8 Tables</td><td><span class="badge badge-pass">TC-SQLI-01 to 05 (PASS)</span></td></tr>
    <tr><td>NFR-SEC-03</td><td>XSS Output Sanitization</td><td>View.php & views/*</td><td>Presentation Layer</td><td><span class="badge badge-pass">TC-XSS-01 (PASS)</span></td></tr>
  </tbody>
</table>

<h1>5. Pre-Submission Final Review Checklist</h1>
<div class="callout">
  <div class="callout-title">📌 FINAL AUDIT QUALITY CERTIFICATION</div>
  This pre-submission quality gate was jointly audited by Team Leader Bautista and SCRUM Master Aquino.
  Every single technical, architectural, testing, and documentation requirement for CCS109 has been verified.
</div>

<table class="data-table">
  <thead>
    <tr>
      <th>Evaluation Criteria</th>
      <th>Required Standard</th>
      <th>Actual Implementation Status</th>
      <th>Verified By</th>
    </tr>
  </thead>
  <tbody>
    <tr><td>3NF Database Normalization</td><td>Min 6 tables, no partial or transitive dependencies</td><td>8 Relational Tables (3NF certified in DATABASE_DESIGN.md)</td><td>Bautista (Lead)</td></tr>
    <tr><td>Zero-Dependency Architecture</td><td>Pure PHP + MySQL, no Node/npm or external build steps</td><td>Pure PHP 8.2 MVC runtime with automated demo fallback</td><td>Aquino (Scrum Master)</td></tr>
    <tr><td>Code Syntax & Linting</td><td>Zero fatal errors or unhandled exceptions across all scripts</td><td>All 44 PHP files passed strict php -l syntax linting</td><td>Balbero (QA)</td></tr>
    <tr><td>Software Test Plan Execution</td><td>Formal test cases with verifiable execution evidence</td><td>IEEE 829 test suite: 36 of 36 automated tests passed (100% success)</td><td>Balbero (QA)</td></tr>
    <tr><td>Security Vulnerability Audit</td><td>Zero high-risk SQLi or XSS vulnerabilities</td><td>100% PDO prepared statements, Bcrypt password hashing verified</td><td>Balbero (QA)</td></tr>
    <tr><td>UI/UX Wireframes Deliverable</td><td>Visual wireframe catalog of all system views</td><td>15 Wireframes captured & cataloged in figma_wireframes/</td><td>Capoy & Baltazar</td></tr>
    <tr><td>System Presentation Deck</td><td>Widescreen PowerPoint slides for case study defense</td><td>18-slide professional PPTX presentation with institutional branding</td><td>Aquino (Scrum Master)</td></tr>
  </tbody>
</table>
"""
    html_doc = build_html_wrapper(
        title="Ethical Analysis, Legal Compliance & Final Review",
        deliverable_code="TASK-1C, 1D & 2D (Bautista & Aquino)",
        author_role="Bautista (Team Leader) & Aquino (SCRUM Master)",
        content_html=html_content
    )
    html_path = os.path.join(OUT_DIR, "Ethical_Analysis_and_Final_Review.html")
    pdf_path = os.path.join(OUT_DIR, "Ethical_Analysis_and_Final_Review.pdf")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_doc)
    html_to_pdf(html_path, pdf_path)
    print(f"Generated PDF: {pdf_path}")
    return pdf_path

if __name__ == "__main__":
    generate_docx()
    generate_pdf()
